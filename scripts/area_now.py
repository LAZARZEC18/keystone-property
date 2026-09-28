"""Newer small-area data to bring the 2021 Census measures up to date. Writes data/official/area-now.json, which
build_suburbs.py reads, so a rebuild never depends on these downloads.

  * ABS Personal Income in Australia (Table 1.4, median total income by SA2, 2018-19 to 2022-23)
  * ABS Wage Price Index (all sectors, total hourly rates excluding bonuses) to carry 2022-23 incomes to today
  * ABS modelled estimates of labour force status by SA4 (6291.0.55.001, monthly)

  python3 scripts/area_now.py
"""
import json, os, subprocess, datetime
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(ROOT, 'data', 'raw', 'pia')
OUT = os.path.join(ROOT, 'data', 'official', 'area-now.json')
ABS = 'https://www.abs.gov.au/statistics'
FILES = {
    't1.xlsx': f'{ABS}/labour/earnings-and-working-conditions/personal-income-australia/2022-23/Table%201%20-%20Total%20income%2C%20earners%20and%20summary%20statistics%20by%20geography%2C%202018-19%20to%202022-23.xlsx',
    'wpi.xlsx': f'{ABS}/economy/price-indexes-and-inflation/wage-price-index-australia/jun-2026/634501.xlsx',
    'MRM1.xlsx': f'{ABS}/labour/employment-and-unemployment/labour-force-australia-detailed/mar-2026/MRM1.xlsx',
}


def fetch():
    os.makedirs(RAW, exist_ok=True)
    for name, url in FILES.items():
        p = os.path.join(RAW, name)
        if not os.path.exists(p):
            subprocess.run(['curl', '-sfL', '-m', '300', '-o', p, url], check=True)


def income():
    ws = openpyxl.load_workbook(os.path.join(RAW, 't1.xlsx'), read_only=True)['Table 1.4']
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[6]
    # Median ($) block is columns 17..21 for 2018-19..2022-23
    out = {}
    for r in rows[7:]:
        code = str(r[0] or '').strip()
        if not (code.isdigit() and len(code) == 9):
            continue
        med = r[17:22]
        if not all(isinstance(v, (int, float)) and v > 0 for v in med):
            continue
        out[code] = {'m1819': med[0], 'm2021': med[2], 'm2122': med[3], 'm2223': med[4]}
    assert hdr[17] == '2018-19' and hdr[21] == '2022-23', hdr[17:22]
    return out


def wpi():
    ws = openpyxl.load_workbook(os.path.join(RAW, 'wpi.xlsx'), read_only=True)['Data1']
    rows = [r for r in ws.iter_rows(values_only=True) if isinstance(r[0], datetime.datetime)]
    s = {r[0].strftime('%Y-%m'): r[3] for r in rows}  # column 3: private and public, all industries
    fy23 = sum(s[k] for k in ('2022-09', '2022-12', '2023-03', '2023-06')) / 4
    last = max(s)
    return s[last] / fy23, last


def unemployment():
    wb = openpyxl.load_workbook(os.path.join(RAW, 'MRM1.xlsx'), read_only=True)
    def table(n):
        rows = list(wb[f'Table {n}'].iter_rows(values_only=True))
        dates = rows[4][1:]
        d = {}
        for r in rows[5:]:
            if not r[0] or not str(r[0])[:3].isdigit():
                continue
            d[str(r[0])[:3]] = {dt.strftime('%Y-%m'): v for dt, v in zip(dates, r[1:]) if isinstance(dt, datetime.datetime) and isinstance(v, (int, float))}
        return d
    emp, une = table(1), table(2)
    out, last = {}, None
    for sa4 in emp:
        e, u = emp[sa4], une.get(sa4, {})
        months = sorted(m for m in e if m in u)
        if len(months) < 60:
            continue
        rate = lambda ms: sum(u[m] for m in ms) / sum(u[m] + e[m] for m in ms) * 100
        recent = months[-12:]
        census = [m for m in ('2021-07', '2021-08', '2021-09') if m in u]
        out[sa4] = {'now': round(rate(recent), 2), 'aug21': round(rate(census), 2)}
        last = months[-1]
    return out, last


def main():
    fetch()
    inc = income()
    f_wpi, wpi_q = wpi()
    une, une_last = unemployment()
    data = {
        'built': datetime.date.today().isoformat(),
        'income': {k: round(v['m2223'] / ((v['m2021'] + v['m2122']) / 2), 4) for k, v in inc.items()},
        'incomeGrowth': {k: round((v['m2223'] / v['m1819'] - 1) * 100, 1) for k, v in inc.items()},
        'wpi': round(f_wpi, 4), 'wpiTo': wpi_q,
        'unemployment': une, 'unemploymentTo': une_last,
        'labels': {
            'income': f'2021 Census household income, moved to {wpi_q[:4]} by the ABS personal income change for the area (to 2022-23) and national wage growth since',
            'incomeGrowth': 'Change in median personal income, 2018-19 to 2022-23 (ABS, by SA2)',
            'unemployment': f'2021 Census rate, moved by the change in the ABS modelled rate for the region (12 months to {une_last})',
        },
    }
    with open(OUT, 'w') as fh:
        json.dump(data, fh, separators=(',', ':'))
    print('area-now:', len(data['income']), 'SA2 incomes;', len(une), 'SA4 labour markets; WPI x', data['wpi'], 'to', wpi_q, '; unemployment to', une_last)


if __name__ == '__main__':
    main()
